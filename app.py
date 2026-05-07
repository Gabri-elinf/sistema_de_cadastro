import tkinter as tk

#Importa componentes adicionais: Label, Button, Entry, TreeView
#Cria caixas de diálogo para escolher e salvar.
from tkinter import ttk, messagebox, filedialog

#Criar, manipular e salvar arquivos Excel.
from openpyxl import Workbook

#Converte números de colunas em letras, utilizado para ajustar as larguras das colunas.
from openpyxl.utils import get_column_letter

from db import conectar

from repo import PessoaRepo

from style import aplicar_estilo, aplicar_zebra_treeview, preencher_treeview

def validar(nome: str, email: str, telefone: str):

    if not nome or len(nome.strip()) < 3:
        return False, "Informe um nome com pelo menos 3 caracteres"

    if not email or '@' not in email:

        return False, "Informe um e-mail válido ( ex.: nome@dominio.com)."

    if not telefone or len(telefone.strip()) < 8:

        return False, "Informe um telefone válido."

    return True, ""


class AppCadastro(tk.Tk):

    def __init__(self):

        super().__init__()

        self.title("Cadastro de Pessoas")

        self.geometry("980x580")

        self.minsize(900, 520)

        aplicar_estilo(self)

        self.conn = conectar(usar_banco=True)

        self.repo = PessoaRepo(self.conn)

        self.var_id = tk.StringVar()

        self.var_nome = tk.StringVar()

        self.var_email = tk.StringVar()

        self.var_telefone = tk.StringVar()

        self.var_pesquisa = tk.StringVar()

        self.var_status = tk.StringVar(value="Pronto.")

        self._montar_formulario()

        self._montar_botoes()

        self._montar_tabela()

        self._montar_statusbar()

        self._centralizar_janela()

        self.bind("<Control-n>", lambda e: self._limpar_campos())

        self.bind("<F5>", lambda e: self.listar_todos())

        self.listar_todos()

        self.protocol("WM_DELETE_WINDOW", self._fechar)


    def _centralizar_janela(self):

        self.update_idletasks()

        largura = self.winfo_width()

        altura = self.winfo_height()

        tela_l = self.winfo_screenwidth()

        tela_a = self.winfo_screenheight()

        x = (tela_l- largura) // 2

        y = (tela_a- altura) // 3

        self.geometry(f"{largura}x{altura}+{x}+{y}")

    def _montar_formulario(self):

        frm = ttk.LabelFrame(self, text = "Dados da Pessoa")

        frm.pack(side=tk.TOP, fill=tk.X, padx = 10, pady = 10)

        ttk.Label(frm, text = "ID:").grid(row = 0, column = 0, padx = 10, pady = 10, stick="e")

        ttk.Entry(frm, textvariable = self.var_id, width=8, state="readonly").grid(
            row = 0, column = 1, padx = 5, pady = 5, stick="w")

        ttk.Label(frm, text="Nome:").grid(row=0, column=2, padx=5, pady=5, stick="e")

        ttk.Entry(frm, textvariable=self.var_nome, width=40).grid(row=0, column=3, padx=5, pady=5, stick="w")

        ttk.Label(frm, text="E-mail:").grid(row=1, column=0, padx=5, pady=5, stick="e")

        ttk.Entry(frm, textvariable=self.var_email, width=40).grid(row=1, column=1, padx=5, pady=5, stick="w")

        ttk.Label(frm, text="Telefone:").grid(row=1, column=2, padx=5, pady=5, stick="e")

        ttk.Entry(frm, textvariable=self.var_telefone, width=22).grid(row=1, column=3, padx=5, pady=5, stick="w")

        ttk.Label(frm, text="Pesquisar (nome, e-mail ou telefone):").grid(row=2, column=0, padx=5, pady=5, stick="e")

        ttk.Entry(frm, textvariable=self.var_pesquisa, width=40).grid(row=2, column=1, columnspan=3, padx=5, pady=5, stick="w")

        frm.grid_columnconfigure(3, weight=1)

    def _montar_botoes(self):
        frm = ttk.Frame(self)
        frm.pack(side=tk.TOP, fill=tk.X, padx=10, pady=5)

        ttk.Button(
            frm,
            text="Cadastrar",
            command=self.cadastrar,
            style="Success.TButton"
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm,
            text="Atualizar",
            command=self.atualizar,
            style="Primary.TButton"
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm,
            text="Excluir",
            command=self.excluir,
            style="Danger.TButton"
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm,
            text="Limpar Campos (Ctrl+N)",
            command=self._limpar_campos,
            style="Secondary.TButton"
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm,
            text="Pesquisar",
            command=self.pesquisar,
            style="Primary.TButton"
        ).pack(side=tk.LEFT, padx=5)

        ttk.Button(
            frm,
            text="Exportar para Excel",
            command=self.exportar_excel,
            style="Excel.TButton"
        ).pack(side=tk.LEFT, padx=5)


    def _montar_tabela(self):

        frm = ttk.LabelFrame(self, text="Registros")

        frm.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx = 10, pady = 5)

        cols = ("id", "nome", "email", "telefone", "criado_em")

        self.tree = ttk.Treeview(frm, columns=cols, show="headings", selectmode="browse")

        self.tree.heading("id", text="ID", command=lambda: self._ordernar_por(0))

        self.tree.heading("nome", text="Nome", command=lambda: self._ordernar_por(1))

        self.tree.heading("email", text="E-mail", command=lambda: self._ordernar_por(2))

        self.tree.heading("telefone", text="Telefone", command=lambda: self._ordernar_por(3))

        self.tree.heading("criado_em", text="Criado em", command=lambda: self._ordernar_por(4))

        self.tree.column("id", width=60, anchor="center")

        self.tree.column("nome", width=260, anchor="w")

        self.tree.column("email", width=260, anchor="w")

        self.tree.column("telefone", width=140, anchor="center")

        self.tree.column("criado_em", width=160, anchor="center")

        vsb = ttk.Scrollbar(frm, orient="vertical", command=self.tree.yview)

        hsb = ttk.Scrollbar(frm, orient="horizontal", command=self.tree.xview)

        #self.tree.configure(yscroll=vsb.set, xscroll = hsb.set)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        vsb.pack(side=tk.RIGHT, fill=tk.Y)

        hsb.pack(side=tk.BOTTOM, fill=tk.X)

        self.tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        aplicar_zebra_treeview(self.tree)

        self.tree.bind("<Double-1>", self._on_duplo_clique)


    def _montar_statusbar(self):

        ttk.Label(self, textvariable=self.var_status, anchor="w").pack(
            side=tk.BOTTOM, fill=tk.X)

    def cadastrar(self):

        nome = self.var_nome.get().strip()

        email = self.var_email.get().strip()

        telefone = self.var_telefone.get().strip()

        ok, msg = validar(nome, email, telefone)

        if not ok:

            messagebox.showwarning("Validação", msg)

            return
        try:

            self.repo.inserir(nome, email, telefone)

            self.var_status.set("Cadastrado realizado com sucesso.")

            self.listar_todos()

            self._limpar_campos()

        except Exception as e:

            messagebox.showerror("Erro ao Cadastrar", f"Falha ao inserir:\n\n{e}")

    def atualizar(self):

        if not self.var_id.get():
            messagebox.showinfo("Atualizar",
                                "Selecione um registro (duplo clique na tabela).")

            return

        nome = self.var_nome.get().strip()

        email = self.var_email.get().strip()

        telefone = self.var_telefone.get().strip()

        ok,msg = validar(nome, email, telefone)

        if not ok:

            messagebox.showwarning("Validação", msg)

            return

        try:

            self.repo.atualizar(int(self.var_id.get()), nome, email, telefone)

            self.var_status.set("Registro atualizado com sucesso.")

            self.listar_todos()

            self._limpar_campos()

        except Exception as e:

            messagebox.showerror("Erro ao Atualizar", f"Falha ao atualizar:\n\n{e}")


    def excluir(self):

        if not self.var_id.get():
            messagebox.showinfo("Excluir",
                                "Selecione um registro (duplo clique na tabela).")

            return

        if not messagebox.askyesno("Confirmação",
                                   "Deseja realmente excluir este registro?"):

            return

        try:

            self.repo.excluir(int(self.var_id.get()))

            self.var_status.set("Registro excluído com sucesso.")

            self.listar_todos()

            self._limpar_campos()

        except Exception as e:

            messagebox.showerror("Erro ao Excluir", f"Falha ao excluir:\n\n{e}")

    def pesquisar(self):

        termo = self.var_pesquisa.get().strip()

        try:

            if not termo:

                rows = self.listar_todos()

                preencher_treeview(self.tree, rows)

                self.var_status.set(f"Total: {len(rows)} registro(s).")

                return

            rows = self.repo.pesquisar(termo)

            preencher_treeview(self.tree, rows)

            self.var_status.set(f"Pesquisa concluída: {len(rows)} registro(s).")

        except Exception as e:

            messagebox.showerror("Erro na Pesquisa", message=f"Falha ao pesquisar:\n\n{e}")

    def listar_todos(self):

        try:
            rows = self.repo.listar_todos()

            preencher_treeview(self.tree, rows)

            self.var_status.set(f"Total: {len(rows)} registro(s).")

            return rows

        except Exception as e:

            messagebox.showerror(title="Erro na listagem", message=f"Falha ao listar:\n\n{e}")

            return []

    def _on_duplo_clique(self, _event):

        item = self.tree.focus()

        if not item:

            return

        vals = self.tree.item(item, "values")

        if not vals:
            return

        _id, nome, email, telefone, criado_em = vals

        self.var_id.set(_id)

        self.var_nome.set(nome)

        self.var_email.set(email)

        self.var_telefone.set(telefone)

    def _limpar_campos(self):

        self.var_id.set("")

        self.var_nome.set("")

        self.var_email.set("")

        self.var_telefone.set("")

        self.var_pesquisa.set("")

        self.tree.selection_remove(self.tree.selection())

    def _ordernar_por(self, coluna_idx):

        dados = [self.tree.item(i, "values") for i in self.tree.get_children()]

        if not dados:

            return

        asc = getattr(self, "_ord_asc", True)

        self._ord_asc = not asc

        def chave(row):

            v = row[coluna_idx]

            if coluna_idx == 0:

                try:

                    return int(v)

                except:

                    return 0

            return str(v).lower()

        ordenado = sorted(dados, key=chave, reverse=asc is False)

        preencher_treeview(self.tree, ordenado)

    def exportar_excel(self):

        linhas = [self.tree.item(i, "values") for i in self.tree.get_children()]

        if not linhas:

            messagebox.showinfo(title="Exportar para Excel", message="Nada para exportar.")

            return

        caminho = filedialog.asksaveasfilename(

            defaultextension="xlsx",

            filetypes=[("Arquivo Excel", "*.xlsx")],

            title="Salvar como"

        )

        if not caminho:

            return

        try:

            wb = Workbook()

            ws = wb.active

            ws.title = "Pessoas"

            cab = ["ID", "Nome", "Email", "Telefone", "Criado em"]

            ws.append(cab)

            for row in linhas:

                ws.append(list(row))

                larguras = [6, 32, 36, 18, 22]

                for idx, largura in enumerate(larguras, start=1):

                    ws.column_dimensions[get_column_letter(idx)].width = largura

            wb.save(caminho)

            messagebox.showinfo(title="Exportar para Excel", message=f"Arquivo salvo com sucesso em:\n {caminho}")

            self.var_status.set(f"Exportado: {caminho}")

        except Exception as e:

            messagebox.showerror(title="Erro ao exportar", message=f"Falha ao exportar:\n\n{e}")


    def _fechar(self):

        try:

            if self.conn:

                self.conn.close()

        except:
            pass

        self.destroy()




















































































































